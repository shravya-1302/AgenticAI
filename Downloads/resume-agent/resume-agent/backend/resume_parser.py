"""Extracts plain text -- and, separately, any real hyperlinks -- from an
uploaded resume file, whatever format it's in.

Why this exists as its own concern: pdfplumber's extract_text() (and
python-docx's paragraph text) only return the *visible characters*. A PDF
made in Word/Canva/etc. almost always has "LinkedIn" and "GitHub" as
clickable hyperlink annotations sitting on top of that plain text, and that
annotation is silently dropped by plain text extraction. Left unhandled,
the tailored resume ends up with "LinkedIn" and "GitHub" as dead, unlinked
words -- exactly the bug this module exists to fix. Every URL surfaced here
comes from something the candidate's own document actually contains
(an embedded link, or a URL typed directly into the text); nothing is ever
guessed or invented.
"""
import io
import re

LINK_LABEL_MATCHERS = [
    ("linkedin", re.compile(r"linkedin\.com", re.I)),
    ("github", re.compile(r"github\.com", re.I)),
    ("twitter", re.compile(r"twitter\.com|x\.com", re.I)),
    ("behance", re.compile(r"behance\.net", re.I)),
    ("medium", re.compile(r"medium\.com", re.I)),
    ("portfolio", re.compile(r"", re.I)),  # never auto-matched; label only
]

URL_IN_TEXT_PATTERN = re.compile(r"https?://[^\s,|]+", re.I)


def _classify_url(url: str) -> str | None:
    if url.lower().startswith("mailto:"):
        return "email"
    for label, pattern in LINK_LABEL_MATCHERS:
        if pattern.pattern and pattern.search(url):
            return label
    return None


def extract_text_from_upload(filename: str, file_bytes: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _extract_pdf(file_bytes)
    if lower.endswith(".docx"):
        return _extract_docx(file_bytes)
    # fall back to treating it as plain text
    return file_bytes.decode("utf-8", errors="ignore")


def extract_links_from_upload(filename: str, file_bytes: bytes) -> dict:
    """Best-effort extraction of real hyperlink URLs, keyed by a normalized
    label (linkedin/github/twitter/behance/medium/email). Returns {} if the
    format has no link extraction path (still safe -- pdf_render simply
    won't linkify anything it has no URL for, rather than fabricating one)."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return _extract_pdf_links(file_bytes)
    if lower.endswith(".docx"):
        return _extract_docx_links(file_bytes)
    return {}


def extract_links_from_text(text: str) -> dict:
    """Fallback for pasted (not uploaded) resume text: pick up any raw URLs
    typed directly into the text itself."""
    links = {}
    for match in URL_IN_TEXT_PATTERN.finditer(text or ""):
        url = match.group(0).rstrip(").,;")
        label = _classify_url(url)
        if label and label not in links:
            links[label] = url
    return links


def _extract_pdf(file_bytes: bytes) -> str:
    import pdfplumber
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
    return "\n".join(text_parts)


def _extract_pdf_links(file_bytes: bytes) -> dict:
    import pdfplumber
    links = {}
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            for hyperlink in page.hyperlinks:
                uri = hyperlink.get("uri")
                if not uri:
                    continue
                label = _classify_url(uri)
                if label and label not in links:
                    links[label] = uri
    return links


def _extract_docx(file_bytes: bytes) -> str:
    import docx
    document = docx.Document(io.BytesIO(file_bytes))
    return "\n".join(p.text for p in document.paragraphs)


def _extract_docx_links(file_bytes: bytes) -> dict:
    import docx
    links = {}
    document = docx.Document(io.BytesIO(file_bytes))
    try:
        rels = document.part.rels
    except AttributeError:
        return links
    for rel in rels.values():
        if rel.reltype.endswith("/hyperlink") and rel.target_ref:
            uri = rel.target_ref
            label = _classify_url(uri)
            if label and label not in links:
                links[label] = uri
    return links


def guess_name_and_contact(resume_text: str) -> tuple[str, str]:
    """Best-effort guess at name (first substantive line) and contact line (has @ or a phone-like pattern)."""
    lines = [l.strip() for l in resume_text.splitlines() if l.strip()]
    name = lines[0][:60] if lines else "Candidate"
    contact = ""
    for l in lines[:8]:
        if "@" in l or re.search(r"\d{3}[-.\s]?\d{3}[-.\s]?\d{4}", l):
            contact = l[:120]
            break
    return name, contact